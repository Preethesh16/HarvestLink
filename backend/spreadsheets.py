"""Read bounded, values-only Excel workbooks; normal CSV validation remains authoritative."""

import base64
import csv
import io
import zipfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADERS = {
    "harvest": [
        "crop",
        "variety",
        "grade",
        "quantity",
        "unit",
        "ask_price",
        "harvest_date",
        "available_date",
    ],
    "inventory": ["crop", "quantity", "unit", "usable_until"],
    "consumption": ["date", "crop", "quantity", "waste", "stockout"],
}


def preview(content: str, kind: str):
    if kind not in HEADERS:
        raise ValueError("Choose harvest, inventory or consumption")
    if len(content) > 2_800_000:
        raise ValueError("Workbook must be smaller than 2 MB")
    try:
        raw = base64.b64decode(content, validate=True)
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist()) > 20_000_000 or len(z.infolist()) > 100:
                raise ValueError("Workbook is too large when expanded")
            if any("vbaProject" in x.filename for x in z.infolist()):
                raise ValueError("Macro-enabled workbooks are not supported")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            ws = wb.active
            if ws.max_row and ws.max_row > 1001:
                raise ValueError("Use at most 1000 data rows")
            if ws.max_column and ws.max_column > len(HEADERS[kind]):
                raise ValueError("Unexpected extra column; use the template headers")
            rows = []
            for cells in ws.iter_rows(max_row=1002, max_col=len(HEADERS[kind]) + 1):
                values = []
                for c in cells:
                    if c.data_type == "f":
                        raise ValueError("Use values only; formulas are not supported")
                    v = c.value
                    if isinstance(v, (date, datetime)):
                        v = v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
                    if v is not None and len(str(v)) > 500:
                        raise ValueError("A cell is too long")
                    values.append("" if v is None else v)
                if not any(v != "" for v in values):
                    continue
                if values[-1] != "":
                    raise ValueError("Unexpected extra column; use the template headers")
                rows.append(values[:-1])
                if len(rows) > 1001:
                    raise ValueError("Use at most 1000 data rows")
            if not rows or rows[0] != HEADERS[kind]:
                raise ValueError("Column headers must match the downloaded template")
            if len(rows) == 1:
                raise ValueError("Workbook has no data rows")
            out = io.StringIO()
            csv.writer(out).writerows(rows)
            return {
                "headers": rows[0],
                "rows": rows[1:6],
                "count": len(rows) - 1,
                "csv": out.getvalue(),
            }
        finally:
            wb.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Cannot read this workbook. Upload a valid .xlsx file.") from exc


def template(kind: str):
    if kind not in HEADERS:
        raise ValueError("Unknown workbook template")
    from datetime import timedelta

    today = date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = kind.title()
    ws.append(HEADERS[kind])
    if kind == "harvest":
        for crop, qty, price in [
            ("Tomato", 100, 30),
            ("Cucumber", 25, 32),
            ("Okra", 30, 40),
            ("Brinjal", 20, 35),
        ]:
            ws.append([crop, "Local", "A", qty, "kg", price, today.isoformat(), today.isoformat()])
    elif kind == "inventory":
        ws.append(["Tomato", 20, "kg", (today + timedelta(days=3)).isoformat()])
    else:
        for i in range(14):
            ws.append([(today - timedelta(days=14 - i)).isoformat(), "Tomato", 8 + i % 3, 0, False])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="234C3B")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 21
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
